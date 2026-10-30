"""Emitter - inventory.xlsx, the full lists behind the brief, for the ADS appendix."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Calibri", size=11)


def _sheet(wb, title, rows, columns=None):
    ws = wb.create_sheet(title[:31])
    if not rows:
        ws.append(["No rows"])
        return
    cols = columns or [k for k in rows[0].keys() if not k.startswith("_")]
    ws.append(cols)
    for r in rows:
        vals = []
        for c in cols:
            v = r.get(c)
            if isinstance(v, (list, tuple, set)):
                v = ", ".join(map(str, v))
            elif isinstance(v, dict):
                v = str(v)
            vals.append(v)
        ws.append(vals)
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = BODY_FONT
    for i, c in enumerate(cols, 1):
        width = max([len(str(c))] + [len(str(r.get(c) or "")) for r in rows[:500]])
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 60)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def write(e, path):
    wb = Workbook()
    wb.remove(wb.active)
    nets = e.networks()
    _sheet(wb, "Subscriptions", e.overview())
    _sheet(wb, "Resources", e.inventory())
    _sheet(wb, "VNets", [{k: v for k, v in n.items() if k not in ("subnets", "peerings")} for n in nets])
    _sheet(wb, "Subnets", [{"vnet": n["name"], "subscription": n["subscription"], **s} for n in nets for s in n["subnets"]])
    _sheet(wb, "Peerings", [{"vnet": n["name"], **p} for n in nets for p in n["peerings"]])
    gw = e.gateways()
    _sheet(wb, "Gateways", gw["gateways"] + gw["connections"] + gw["localGateways"] + gw["circuits"])
    _sheet(wb, "EdgeServices", e.edge_services())
    _sheet(wb, "RouteTables", e.route_tables())
    _sheet(wb, "PrivateEndpoints", e.private_endpoints())
    d = e.dns()
    _sheet(wb, "PrivateDNS", d["private"])
    _sheet(wb, "PublicIPs", e.public_ips())
    _sheet(wb, "OpenNSGRules", e.open_nsg_rules())
    _sheet(wb, "PaaSNetworkAccess", e.paas_network_access())
    _sheet(wb, "VirtualMachines", e.vms())
    asv = e.app_service()
    _sheet(wb, "AppServicePlans", asv["plans"])
    _sheet(wb, "Apps", asv["apps"])
    _sheet(wb, "Data", e.data())
    _sheet(wb, "Access", e.access())
    _sheet(wb, "KeyVaults", e.key_vaults())
    _sheet(wb, "Policy", e.policy())
    _sheet(wb, "Backup", e.backup_items)
    _sheet(wb, "Advisor", e.advisor_recs)
    _sheet(wb, "SecurityFindings", e.findings)
    _sheet(wb, "Costs", [{"subscription": e.sub(c["subscriptionId"]), **c} for c in e.cost_rows])
    sp = e.spend()
    if sp["months"]:
        subs = sorted({k for m in sp["months"] for k in m["bySub"]})
        _sheet(wb, "CostByMonth", [{"month": m["month"], "daysCovered": m["days"],
                                    "wholeMonth": "yes" if m["complete"] else "no",
                                    **{s: m["bySub"].get(s) for s in subs},
                                    "total": m["cost"], "changeOnPriorMonth": m["change"]} for m in sp["months"]])
    _sheet(wb, "Dependencies", e.dependencies())
    _sheet(wb, "OutsideScope", e.scope_gaps()["references"])
    _sheet(wb, "Retirements", [{k: v for k, v in r.items() if not k.startswith("_")} for r in e.retirements()])
    _sheet(wb, "Billing", e.billing()["subscriptions"])
    # Sheets below exist only when the matching raw data was collected.
    if e.collected("site_config"):
        _sheet(wb, "SiteConfig", e.site_facts())
    if e.collected("cosmos_config"):
        cz = e.cosmos_config()
        _sheet(wb, "CosmosAccounts", cz["accounts"])
        _sheet(wb, "CosmosContainers", cz["containers"])
    if e.collected("servicebus_config"):
        sb = e.servicebus_config()
        _sheet(wb, "ServiceBus", [{k: v for k, v in n.items() if k != "topWithoutSubscriptions"} for n in sb["namespaces"]])
        _sheet(wb, "ServiceBusEntities", sb["entities"])
    ds = e.diagnostic_settings()
    if ds:
        _sheet(wb, "DiagnosticSettings", ds["settings"])
        _sheet(wb, "DiagnosticCoverage", ds["coverage"])
    ml = e.measured_load()
    if ml:
        _sheet(wb, "MeasuredCosmos", ml["cosmos"])
        _sheet(wb, "MeasuredCosmosContainers", ml["containers"])
        _sheet(wb, "MeasuredServiceBus", ml["servicebus"])
        _sheet(wb, "MeasuredPlans", ml["plans"])
        _sheet(wb, "MeasuredSQL", ml["sql"])
    lau = e.log_analytics_usage()
    if lau:
        _sheet(wb, "LogAnalyticsUsage", [{**{k: v for k, v in u.items() if k not in ("topTables", "overrides")},
                                          "largestTables": "; ".join(f"{t['table']} {t['gb']} GB" for t in u["topTables"]),
                                          "retentionOverrides": "; ".join(u["overrides"])} for u in lau])
    en = e.entra()
    if en:
        _sheet(wb, "ConditionalAccess", en["policies"])
        _sheet(wb, "DirectoryRoles", en["roles"])
    wb.save(path)
